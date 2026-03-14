"""
Document parser for extracting text from various file formats.
"""
from pathlib import Path
from typing import List


class DocumentParser:
    SUPPORTED_EXTENSIONS = {
        '.txt', '.md', '.csv', '.json', '.py', '.js', '.html', '.css', '.xml',
        '.pdf', '.docx', '.xlsx', '.pptx'
    }
    
    @staticmethod
    def parse(file_path: str) -> str:
        """
        Parse a document and extract text content.
        
        Args:
            file_path: Path to the document file
            
        Returns:
            Extracted text content
            
        Raises:
            ValueError: If file type is not supported
            FileNotFoundError: If file doesn't exist
        """
        path = Path(file_path)
        
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        extension = path.suffix.lower()
        
        if extension not in DocumentParser.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: {extension}")
        
        if extension in {'.txt', '.md', '.csv', '.json', '.py', '.js', '.html', '.css', '.xml'}:
            return DocumentParser._parse_text(file_path)
        elif extension == '.pdf':
            return DocumentParser._parse_pdf(file_path)
        elif extension == '.docx':
            return DocumentParser._parse_docx(file_path)
        elif extension == '.xlsx':
            return DocumentParser._parse_xlsx(file_path)
        elif extension == '.pptx':
            return DocumentParser._parse_pptx(file_path)
        
        raise ValueError(f"Parser not implemented for: {extension}")
    
    @staticmethod
    def _parse_text(file_path: str) -> str:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    
    @staticmethod
    def _parse_pdf(file_path: str) -> str:
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ImportError("pypdf is required for PDF parsing. Install with: pip install pypdf")
        
        reader = PdfReader(file_path)
        text_parts = []
        
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        
        return '\n\n'.join(text_parts)
    
    @staticmethod
    def _parse_docx(file_path: str) -> str:
        try:
            from docx import Document
        except ImportError:
            raise ImportError("python-docx is required for DOCX parsing. Install with: pip install python-docx")
        
        doc = Document(file_path)
        paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
        
        return '\n\n'.join(paragraphs)
    
    @staticmethod
    def _parse_xlsx(file_path: str) -> str:
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ImportError("openpyxl is required for XLSX parsing. Install with: pip install openpyxl")
        
        wb = load_workbook(file_path, data_only=True)
        text_parts = []
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            sheet_text = [f"Sheet: {sheet_name}"]
            
            for row in sheet.iter_rows(values_only=True):
                row_text = ' | '.join(str(cell) for cell in row if cell is not None)
                if row_text.strip():
                    sheet_text.append(row_text)
            
            if len(sheet_text) > 1:
                text_parts.append('\n'.join(sheet_text))
        
        return '\n\n'.join(text_parts)
    
    @staticmethod
    def _parse_pptx(file_path: str) -> str:
        try:
            from pptx import Presentation
        except ImportError:
            raise ImportError("python-pptx is required for PPTX parsing. Install with: pip install python-pptx")
        
        prs = Presentation(file_path)
        text_parts = []
        
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_text = [f"Slide {slide_num}"]
            
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    shape_text = getattr(shape, "text", "")
                    if shape_text.strip():
                        slide_text.append(shape_text)
            
            if len(slide_text) > 1:
                text_parts.append('\n'.join(slide_text))
        
        return '\n\n'.join(text_parts)
    
    @staticmethod
    def get_file_type(file_path: str) -> str:
        extension = Path(file_path).suffix.lower()
        
        type_map = {
            '.txt': 'text',
            '.md': 'text',
            '.csv': 'text',
            '.json': 'text',
            '.pdf': 'pdf',
            '.docx': 'docx',
            '.xlsx': 'xlsx',
            '.pptx': 'pptx'
        }
        
        return type_map.get(extension, 'unknown')
    
    @staticmethod
    def is_supported(file_path: str) -> bool:
        extension = Path(file_path).suffix.lower()
        return extension in DocumentParser.SUPPORTED_EXTENSIONS
